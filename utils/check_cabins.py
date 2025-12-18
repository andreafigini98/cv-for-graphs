import re
import csv
from openpyxl import load_workbook
import math



from openpyxl import load_workbook

def load_competenze_xlsx(path_xlsx, debug=True):
    wb = load_workbook(path_xlsx, data_only=True)

    if debug:
        print("📘 Fogli disponibili:", wb.sheetnames)

    ws = wb.active

    if debug:
        print("📄 Foglio attivo:", ws.title)
        print("📐 Max row:", ws.max_row)

    comp_e_raw = []
    comp_d_raw = []

    for row in range(4, ws.max_row + 1):
        val_g = ws[f"G{row}"].value
        val_s = ws[f"S{row}"].value

        if debug and row < 15:  # stampiamo solo le prime righe
            print(f"Riga {row} | G: {val_g!r} | S: {val_s!r}")

        if isinstance(val_g, str) and val_g.strip():
            comp_e_raw.append(val_g.strip())

        if isinstance(val_s, str) and val_s.strip():
            comp_d_raw.append(val_s.strip())

    if debug:
        print("📥 Raw competenza E (prime 5):", comp_e_raw[:5])
        print("📥 Raw competenza D (prime 5):", comp_d_raw[:5])
        print("📊 Totale raw E:", len(comp_e_raw))
        print("📊 Totale raw D:", len(comp_d_raw))

    # Normalizzazione
    comp_e_norm = set()
    comp_d_norm = set()

    for s in comp_e_raw:
        norm = normalize_single_cabin_id(s)
        if norm:
            comp_e_norm.add(norm)
        elif debug:
            print("⚠️ Codice E non riconosciuto:", repr(s))

    for s in comp_d_raw:
        norm = normalize_single_cabin_id(s)
        if norm:
            comp_d_norm.add(norm)
        elif debug:
            print("⚠️ Codice D non riconosciuto:", repr(s))

    if debug:
        print("✅ Competenza E normalizzata (prime 5):", list(comp_e_norm)[:5])
        print("✅ Competenza D normalizzata (prime 5):", list(comp_d_norm)[:5])

    return comp_e_norm, comp_d_norm




def load_xlsx_cabins(path_xlsx):
    wb = load_workbook(path_xlsx, data_only=True)
    ws = wb.active

    cabins = []

    # seconda colonna = B
    # righe pari a partire dalla 10
    for row in range(10, ws.max_row + 1, 2):
        val = ws.cell(row=row, column=2).value

        if val is None:
            continue

        # scarta NaN
        if isinstance(val, float) and math.isnan(val):
            continue

        cabins.append(str(val).strip())

    return cabins










def normalize_single_cabin_id(s: str) -> str:
    """
    Normalizza un codice cabina in formato canonico:
    AA 10-2-262241

    Gestisce:
    - DU40-2-453048
    - DU102-134210_TR01
    - DU102-134210
    - DU40.2.453048
    """
    if not isinstance(s, str):
        return ""

    s = s.upper().strip()

    # 1️⃣ pulizia base
    s = s.replace(".", "-")
    s = re.sub(r"_?TR\d+", "", s)      # rimuove _TR01, TR02, ecc
    s = re.sub(r"\s+", " ", s)

    # 2️⃣ FORMATO A: DU40-2-453048
    m = re.search(r"([A-Z]{2})\s*(\d{2})-(\d)-(\d{6})", s)
    if m:
        return f"{m.group(1)} {m.group(2)}-{m.group(3)}-{m.group(4)}"

    # 3️⃣ FORMATO B: DU102-134210  → 10-2
    m = re.search(r"([A-Z]{2})\s*(\d{3})-(\d{6})", s)
    if m:
        aa = m.group(1)
        zone = m.group(2)[:2]   # prime 2 cifre
        sub = m.group(2)[2]     # terza cifra
        code = m.group(3)
        return f"{aa} {zone}-{sub}-{code}"

    # 4️⃣ non riconosciuto
    return ""








def normalize_cabin_id(xlsx_cabin_list):
    """
    Input: lista codici XLSX
    Output: set di codici normalizzati stile 'AA 10-2-262241'
    """
    normalized = set()

    for s in xlsx_cabin_list:
        norm = normalize_single_cabin_id(s)

        if norm:
            normalized.add(norm)
        else:
            print(f"⚠️ Codice XLSX non riconosciuto: {s}")

    return normalized





def assign_competenze(associations, xlsx_cabin_set, set_comp_e, set_comp_d, debug=True):

    new_cabin_set = normalize_cabin_id(xlsx_cabin_set)

    for i, a in enumerate(associations):

        a["in_xlsx"] = a["parsed_id"] in new_cabin_set

        if(a["in_xlsx"] == False):
            raw_id = a.get("parsed_id")
            #norm_id = normalize_single_cabin_id(raw_id)

            in_e = raw_id in set_comp_e
            in_d = raw_id in set_comp_d

            a["competenza_e"] = in_e
            a["competenza_d"] = in_d

            # 🔍 DEBUG sulle prime N
            if debug and i < 15:
                print("———")
                print(f"Cabina idx {a['cabina_index']}")
                print(" raw_id :", repr(raw_id))
                print(" in E   :", in_e)
                print(" in D   :", in_d)


    return associations





def write_csv(associations, csv_out: str = "associations.csv"):
    ##### OLD #####
    with open("outputs/associations_old.csv", "w", newline="", encoding="utf-8") as f:
        #writer = csv.writer(f, quoting=csv.QUOTE_NONE, escapechar='\\')
        writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_NONE, escapechar='\\')

        writer.writerow(["cabina_index", "cabina_id", "cabina_conf", "info_text"])
        for a in associations:
            info_text = a["info_text"].replace("\n", " ").replace("\r", "").strip().strip('"')
            writer.writerow([a["cabina_index"], a["cabina_id"], f"{a['cabina_conf']:.1f}", info_text])



    ##### NEW #####

    with open(csv_out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(
            f,
            delimiter=";",
            quoting=csv.QUOTE_NONE,
            escapechar="\\"
        )

        writer.writerow([
            "cabina_index",
            "sigla_cabina",
            "cabina_id",
            "info_text",
            "trasformatore",
            "utenza",
            "gruppo",
            "in_xlsx",
            "competenza_e",
            "competenza_d",
        ])

        for a in associations:

            writer.writerow([
                a["cabina_index"],
                a["cabina_id"],   # ← la sigla (es. MB)
                a["parsed_id"],        # ← DU 10-2-xxxxxx
                a["info_text_clean"],
                a["trasformatore"],
                a["utenza"],
                a["gruppo"],
                #a["parsed_id"] in new_cabin_set,  
                a["in_xlsx"],
                a.get("competenza_e", ""),
                a.get("competenza_d", ""),
            ])






from openpyxl import Workbook
from openpyxl.styles import PatternFill

def write_xlsx_colored(associations, out_xlsx):
    wb = Workbook()
    ws = wb.active
    ws.title = "associations"

    # colori tenui
    FILL_TRUE  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    FILL_FALSE = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")

    headers = [
        "cabina_index",
        "sigla_cabina",
        "cabina_id",
        "info_text",
        "trasformatore",
        "utenza",
        "gruppo",
        "in_xlsx",
        "competenza_e",
        "competenza_d",
    ]

    ws.append(headers)

    for a in associations:
        row = [
            a.get("cabina_index"),
            a.get("cabina_id"),
            a.get("parsed_id"),
            a.get("info_text_clean", ""),
            a.get("trasformatore", ""),
            a.get("utenza", ""),
            a.get("gruppo", ""),
            a.get("in_xlsx"),
            a.get("competenza_e"),
            a.get("competenza_d"),
        ]

        ws.append(row)
        r = ws.max_row

        # colonne booleane (1-based)
        bool_cols = {
            8: a.get("in_xlsx"),
            9: a.get("competenza_e"),
            10: a.get("competenza_d"),
        }

        for col_idx, val in bool_cols.items():
            cell = ws.cell(row=r, column=col_idx)
            if val is True:
                cell.fill = FILL_TRUE
            elif val is False:
                cell.fill = FILL_FALSE
            # None → lasciamo bianco (come richiesto)

    wb.save(out_xlsx)
    print(f"✅ XLSX colorato salvato in {out_xlsx}")
